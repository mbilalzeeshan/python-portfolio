"""Web scraper demo - scrapes the public practice site books.toscrape.com.

Extracts title, price, star rating and availability for every book across
all catalogue pages, then saves clean data to books.csv and books.xlsx.

books.toscrape.com is a public site built for practising scraping - no
login, no paywall. The scraper is polite: a real User-Agent header and a
short delay between page requests.

Run:
    pip install -r requirements.txt
    python scrape_books.py
"""

import csv
import re
import time

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_URL = "http://books.toscrape.com/catalogue/page-{}.html"
HEADERS = {
    "User-Agent": "Mozilla/5.0 (personal project) requests-based scraper"
}
DELAY_SECONDS = 0.7  # polite crawling: small pause between page requests

RATING_WORDS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


def scrape_page(page_number):
    """Return a list of book dicts from one catalogue page."""
    url = BASE_URL.format(page_number)
    response = requests.get(url, headers=HEADERS, timeout=15)
    if response.status_code == 404:
        return None  # past the last page
    response.raise_for_status()
    # The site serves UTF-8 bytes but declares no usable charset, so set it
    # explicitly to avoid mojibake like 'Â£51.77'.
    response.encoding = "utf-8"

    soup = BeautifulSoup(response.text, "html.parser")
    books = []
    for article in soup.select("article.product_pod"):
        title = article.h3.a["title"].strip()
        price_text = article.select_one("p.price_color").get_text(strip=True)
        price = float(re.sub(r"[^\d.]", "", price_text))

        rating_classes = article.select_one("p.star-rating")["class"]
        rating_word = next(c for c in rating_classes if c != "star-rating")
        rating = RATING_WORDS.get(rating_word, 0)

        availability = article.select_one("p.instock.availability").get_text(strip=True)

        books.append({
            "title": title,
            "price_gbp": price,
            "rating_out_of_5": rating,
            "availability": availability,
        })
    return books


def scrape_all():
    """Walk every catalogue page until a 404, pausing between requests."""
    all_books = []
    page = 1
    while True:
        books = scrape_page(page)
        if not books:  # 404 or empty page = no more pages
            break
        all_books.extend(books)
        print(f"Page {page}: {len(books)} books")
        page += 1
        time.sleep(DELAY_SECONDS)
    return all_books


def save_csv(books, path="books.csv"):
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f, fieldnames=["title", "price_gbp", "rating_out_of_5", "availability"]
        )
        writer.writeheader()
        writer.writerows(books)
    print(f"Saved {path} ({len(books)} rows)")


def save_excel(books, path="books.xlsx"):
    wb = Workbook()
    ws = wb.active
    ws.title = "Books"

    headers = ["Title", "Price (GBP)", "Rating (out of 5)", "Availability"]
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    for row, book in enumerate(books, start=2):
        ws.cell(row=row, column=1, value=book["title"])
        price_cell = ws.cell(row=row, column=2, value=book["price_gbp"])
        price_cell.number_format = '\u00a3#,##0.00'
        ws.cell(row=row, column=3, value=book["rating_out_of_5"])
        ws.cell(row=row, column=4, value=book["availability"])

    widths = [60, 14, 18, 22]
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)
    print(f"Saved {path} ({len(books)} rows)")


def main():
    print("Scraping books.toscrape.com ...")
    books = scrape_all()
    print(f"Total books scraped: {len(books)}")
    save_csv(books)
    save_excel(books)


if __name__ == "__main__":
    main()
