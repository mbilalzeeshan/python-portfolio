"""Sales report generator - reads sample_sales.csv and builds a formatted
Excel report with pandas + openpyxl.

Output (monthly_report.xlsx):
  - Summary sheet: revenue totals by month, by product, and grand total,
    plus a bar chart of revenue per product.
  - Data sheet: every sale row with a Revenue column, styled headers,
    tuned column widths, currency and date formats.

Run:
    pip install -r requirements.txt
    python generate_report.py
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference

INPUT_CSV = "sample_sales.csv"
OUTPUT_XLSX = "monthly_report.xlsx"

HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
TITLE_FONT = Font(color="1F4E79", bold=True, size=14)
MONEY_FORMAT = "$#,##0.00"
THIN_BORDER = Border(
    left=Side(style="thin", color="B0B0B0"),
    right=Side(style="thin", color="B0B0B0"),
    top=Side(style="thin", color="B0B0B0"),
    bottom=Side(style="thin", color="B0B0B0"),
)


def style_header_row(ws, max_col, row=1):
    for col in range(1, max_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 22


def build_dataframes():
    df = pd.read_csv(INPUT_CSV, parse_dates=["Date"])
    df["Revenue"] = df["Units"] * df["UnitPrice"]
    df["Month"] = df["Date"].dt.strftime("%Y-%m")

    by_month = df.groupby("Month", as_index=False)["Revenue"].sum()
    by_product = df.groupby("Product", as_index=False)["Revenue"].sum()
    by_product = by_product.sort_values("Revenue", ascending=False)
    grand_total = df["Revenue"].sum()
    return df, by_month, by_product, grand_total


def write_summary_sheet(wb, by_month, by_product, grand_total):
    ws = wb.create_sheet("Summary")

    ws["A1"] = "Monthly Sales Report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = "Generated from sample_sales.csv"
    ws["A2"].font = Font(italic=True, color="666666")

    # Revenue by month
    ws["A4"] = "Revenue by Month"
    ws["A4"].font = Font(bold=True, size=12)
    ws["A5"] = "Month"
    ws["B5"] = "Revenue"
    style_header_row(ws, 2, row=5)
    for i, (_, row) in enumerate(by_month.iterrows(), start=6):
        ws.cell(row=i, column=1, value=row["Month"]).border = THIN_BORDER
        rev = ws.cell(row=i, column=2, value=round(row["Revenue"], 2))
        rev.number_format = MONEY_FORMAT
        rev.border = THIN_BORDER

    # Revenue by product (starts a few rows below the month table)
    prod_start = 6 + len(by_month) + 2
    ws.cell(row=prod_start, column=1, value="Revenue by Product").font = Font(bold=True, size=12)
    header_row = prod_start + 1
    ws.cell(row=header_row, column=1, value="Product")
    ws.cell(row=header_row, column=2, value="Revenue")
    style_header_row(ws, 2, row=header_row)
    for i, (_, row) in enumerate(by_product.iterrows(), start=header_row + 1):
        ws.cell(row=i, column=1, value=row["Product"]).border = THIN_BORDER
        rev = ws.cell(row=i, column=2, value=round(row["Revenue"], 2))
        rev.number_format = MONEY_FORMAT
        rev.border = THIN_BORDER

    total_row = header_row + 1 + len(by_product)
    ws.cell(row=total_row, column=1, value="GRAND TOTAL").font = Font(bold=True)
    grand = ws.cell(row=total_row, column=2, value=round(grand_total, 2))
    grand.number_format = MONEY_FORMAT
    grand.font = Font(bold=True)

    for col, width in (("A", 24), ("B", 16)):
        ws.column_dimensions[col].width = width

    # Bar chart: revenue per product
    chart = BarChart()
    chart.type = "col"
    chart.title = "Revenue by Product"
    chart.y_axis.title = "Revenue ($)"
    data = Reference(ws, min_col=2, min_row=header_row, max_row=header_row + len(by_product))
    cats = Reference(ws, min_col=1, min_row=header_row + 1, max_row=header_row + len(by_product))
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(cats)
    chart.shape = 4
    chart.height = 7.5
    chart.width = 14
    ws.add_chart(chart, "D4")

    return ws


def format_data_sheet(wb, df):
    ws = wb["Data"]
    style_header_row(ws, len(df.columns))

    widths = {"Date": 14, "Product": 22, "Region": 12, "Units": 10,
              "UnitPrice": 12, "Revenue": 14, "Month": 10}
    for idx, col in enumerate(df.columns, start=1):
        letter = get_column_letter(idx)
        ws.column_dimensions[letter].width = widths.get(col, 16)
        for row in range(2, len(df) + 2):
            ws.cell(row=row, column=idx).border = THIN_BORDER
            ws.cell(row=row, column=idx).alignment = Alignment(horizontal="center")

    for row in range(2, len(df) + 2):
        ws.cell(row=row, column=1).number_format = "YYYY-MM-DD"
        ws.cell(row=row, column=5).number_format = MONEY_FORMAT
        ws.cell(row=row, column=6).number_format = MONEY_FORMAT

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    df, by_month, by_product, grand_total = build_dataframes()

    # Write the raw data sheet with pandas, then style everything with openpyxl
    with pd.ExcelWriter(OUTPUT_XLSX, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Data", index=False)

    wb = load_workbook(OUTPUT_XLSX)
    format_data_sheet(wb, df)
    write_summary_sheet(wb, by_month, by_product, grand_total)

    # Put Summary first
    wb.move_sheet("Summary", offset=-1)
    wb.save(OUTPUT_XLSX)

    print(f"Saved {OUTPUT_XLSX}")
    print(f"  Rows: {len(df)} | Months: {len(by_month)} | Products: {len(by_product)}")
    print(f"  Grand total revenue: ${grand_total:,.2f}")


if __name__ == "__main__":
    main()
